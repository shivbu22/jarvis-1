"""Tests for /api/upload Multi-modal File Intake Pipeline."""

import io
from fastapi.testclient import TestClient
from apps.gateway.main import create_app


def test_upload_text_file(tmp_path):
    app = create_app()
    with TestClient(app) as client:
        file_content = b"Alpha Protocol: Autonomous System Evolution and Analysis."
        files = {"file": ("alpha_notes.txt", io.BytesIO(file_content), "text/plain")}

        resp = client.post("/api/upload", files=files)
        assert resp.status_code == 200
        data = resp.json()
        assert data["ok"] is True
        assert data["filename"] == "alpha_notes.txt"
        assert "Alpha Protocol" in data["extracted_text"]
        assert data["is_image"] is False


def test_upload_empty_file_rejected():
    app = create_app()
    with TestClient(app) as client:
        files = {"file": ("empty.txt", io.BytesIO(b""), "text/plain")}
        resp = client.post("/api/upload", files=files)
        assert resp.status_code == 400


def test_upload_docx_file():
    app = create_app()
    with TestClient(app) as client:
        import docx
        doc = docx.Document()
        doc.add_heading("Quantum Execution Plan", level=1)
        doc.add_paragraph("All subsystems fully synchronized.")
        buf = io.BytesIO()
        doc.save(buf)
        buf.seek(0)

        files = {"file": ("plan.docx", buf, "application/vnd.openxmlformats-officedocument.wordprocessingml.document")}
        resp = client.post("/api/upload", files=files)
        assert resp.status_code == 200
        data = resp.json()
        assert data["ok"] is True
        assert data["is_document"] is True
        assert "Quantum Execution Plan" in data["extracted_text"]
